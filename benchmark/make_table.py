#!/usr/bin/env python
"""Turn benchmark results into Supplementary Table S1 (Word + CSV).

python benchmark/make_table.py --results ~/hb_benchmark/results_published \
       [--sensitivity ~/hb_benchmark/results_nobfill] --out ~/hb_benchmark/Table_S1.docx
"""
import argparse
import json
import os
import sys

import numpy as np
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from hbbench import models as M  # noqa: E402

REPR = {"last_hb_binary": "Last Hb value", "last_hb_continuous": "Last Hb value",
        "logreg": "51 engineered features", "rf": "51 engineered features", "xgb": "51 engineered features",
        "xgb_calibrated_published": "51 engineered features", "lstm": "Hourly sequence, 19 h x 19 variables"}


def rows(res_dir):
    test = pd.read_csv(os.path.join(res_dir, "test_metrics.csv"))
    cvp = os.path.join(res_dir, "cv_summary.csv")
    cv = pd.read_csv(cvp).set_index("model") if os.path.exists(cvp) else pd.DataFrame()
    info = json.load(open(os.path.join(res_dir, "run_info.json")))
    out = []
    for n in info["models"]:
        r = {"Model": M.LABELS[n], "Input": REPR[n]}
        if n in cv.index:
            c = cv.loc[n]
            r["CV AUROC, mean ± SD"] = f"{c.auroc_mean:.3f} ± {c.auroc_sd:.3f}"
            r["CV AUPRC, mean ± SD"] = f"{c.auprc_mean:.3f} ± {c.auprc_sd:.3f}"
        else:
            r["CV AUROC, mean ± SD"] = r["CV AUPRC, mean ± SD"] = "–"
        for split, lab in [("internal_test", "Internal test"), ("external_mimic", "MIMIC-IV")]:
            t = test[(test.model == n) & (test.split == split)]
            if len(t):
                t = t.iloc[0]
                r[f"{lab} AUROC (95% CI)"] = f"{t.auroc:.3f} ({t.auroc_lo:.3f}–{t.auroc_hi:.3f})"
                r[f"{lab} AUPRC (95% CI)"] = f"{t.auprc:.3f} ({t.auprc_lo:.3f}–{t.auprc_hi:.3f})"
                if split == "internal_test":
                    r["Internal test Brier"] = "–" if np.isnan(t.brier) else f"{t.brier:.3f}"
        out.append(r)
    return pd.DataFrame(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", required=True)
    ap.add_argument("--sensitivity", default=None, help="results dir of the --no-bfill run")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    res = os.path.expanduser(a.results)
    tab = rows(res)
    out = os.path.expanduser(a.out)
    tab.to_csv(os.path.splitext(out)[0] + ".csv", index=False)
    sens = rows(os.path.expanduser(a.sensitivity)) if a.sensitivity else None
    try:
        import docx
        from docx.enum.section import WD_ORIENT
        from docx.shared import Pt
    except ImportError:
        print("python-docx not installed -> CSV only:", os.path.splitext(out)[0] + ".csv")
        return
    d = docx.Document()
    s = d.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = s.page_height, s.page_width
    d.styles["Normal"].font.name = "Arial"
    d.styles["Normal"].font.size = Pt(8)

    def add(df, title):
        d.add_paragraph().add_run(title).bold = True
        t = d.add_table(rows=1, cols=len(df.columns))
        t.style = "Table Grid"
        for c, h in zip(t.rows[0].cells, df.columns):
            c.text = h
            c.paragraphs[0].runs[0].bold = True
        for _, r in df.iterrows():
            for c, v in zip(t.add_row().cells, r.values):
                c.text = "" if pd.isna(v) else str(v)

    add(tab, "Supplementary Table S1. Comparison of candidate model classes for 6-hour prediction of severe anemia (Hb <8 g/dL)")
    d.add_paragraph(
        "All models were developed on the ViennaAIdb training set (70% of patients) with identical outcome definition, "
        "preprocessing and patient-grouped five-fold cross-validation (GroupKFold on patient ID); the internal test set "
        "(30% of patients) and MIMIC-IV were used only for final evaluation. Tabular models used the same 51 engineered "
        "features; the LSTM used the underlying hourly sequences. Logistic regression, random forest, XGBoost and LSTM were "
        "trained with class-balancing and transition weights and were not recalibrated, so their Brier scores are not "
        "comparable with the calibrated final model. 95% CIs: patient-level bootstrap. AUPRC, area under the "
        "precision-recall curve; AUROC, area under the receiver operating characteristic curve; CV, cross-validation; "
        "Hb, hemoglobin; LSTM, long short-term memory network; SD, standard deviation.")
    if sens is not None:
        add(sens, "Sensitivity analysis: preprocessing without backward filling of predictors")
    d.save(out)
    print("written", out)


if __name__ == "__main__":
    main()
